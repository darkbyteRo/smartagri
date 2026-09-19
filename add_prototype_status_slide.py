"""
Add a 'Prototype Status' note to the last slide of the SIH presentation.
This adds a text box at the bottom of Slide 6 with a clear summary of 
what is ready in the prototype and what remains in progress.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

PPTX_PATH = os.path.join(os.path.dirname(__file__), "test_replaced.pptx")
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "test_replaced.pptx")  # overwrite

prs = Presentation(PPTX_PATH)

# --- Target: Slide 6 (index 5) – "RESEARCH AND REFERENCES" ---
last_slide = prs.slides[len(prs.slides) - 1]

# Add a text box at the bottom-left area of the last slide
# Slide dimensions are widescreen 16:9 (13.333" x 7.5")
# Position: near the bottom, spanning most of the width
left = Emu(173736)        # Align with the research table
top = Emu(5900000)        # Near the bottom, above the footer bar
width = Emu(11000000)     # Wide enough
height = Emu(500000)      # Compact

txBox = last_slide.shapes.add_textbox(left, top, width, height)
tf = txBox.text_frame
tf.word_wrap = True

# --- Main status line ---
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT

run_label = p.add_run()
run_label.text = "⚙ PROTOTYPE STATUS: "
run_label.font.size = Pt(9)
run_label.font.bold = True
run_label.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)  # Blue

run_done = p.add_run()
run_done.text = (
    "✅ Backend (FastAPI + SQLite): Auth, Market Intelligence, Net Realization Engine, "
    "Sell/Hold Engine, Buyer Matching, Price Prediction (ML), Telugu AI Assistant (Sarvam AI), "
    "Listings & Offers, Transaction Flow — ALL WORKING.  "
    "✅ Frontend (Next.js): Farmer Dashboard, Buyer Dashboard, Admin Panel, Market Comparison, "
    "AI Chatbot (Voice + Text) — ALL WORKING.  "
    "✅ Tests: Auth, Buyer Matching, Sell/Hold, Market Intelligence — PASSING.  "
)
run_done.font.size = Pt(8)
run_done.font.color.rgb = RGBColor(0x22, 0x8B, 0x22)  # Green

run_wip = p.add_run()
run_wip.text = (
    "🔄 In Progress: PostgreSQL migration (currently SQLite for demo), "
    "Agmarknet live API integration, PWA offline mode."
)
run_wip.font.size = Pt(8)
run_wip.font.color.rgb = RGBColor(0xFF, 0x8C, 0x00)  # Orange

# Add a second paragraph with the note
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.LEFT
run_note = p2.add_run()
run_note.text = (
    "Note: This is a working prototype with simulated Telangana market data. "
    "Core engines (Net Realization, Sell/Hold, Buyer Matching, Telugu AI) are fully functional end-to-end. "
    "GitHub: github.com/darkbyteRo/smartagri"
)
run_note.font.size = Pt(7)
run_note.font.italic = True
run_note.font.color.rgb = RGBColor(0x55, 0x55, 0x55)  # Gray

prs.save(OUTPUT_PATH)
print(f"DONE: Prototype status note added to last slide. Saved -> {OUTPUT_PATH}")
